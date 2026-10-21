"""Document bytes to text with stable offsets (design 9.2 step 2).

Markdown and plain text are decoded natively (UTF-8, with a BOM tolerated). PDF goes through
`pypdf` and DOCX through `python-docx`, both from the `documents` extra of this package; when
the extra is missing the job fails with a message naming it. OCR is out of scope: a PDF whose
pages hold no text is reported as scanned and yields no text.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

MARKDOWN = "text/markdown"
PLAIN_TEXT = "text/plain"
PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MEDIA_TYPES = (MARKDOWN, PLAIN_TEXT, PDF, DOCX)
EXTRA_HINT = "install lorebase-scribe[documents]"


class ParseError(ValueError):
    """The document cannot be turned into text; retrying will not help."""


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    """Parsed text; `pages` holds the start offset of each page (PDF only)."""

    text: str
    media_type: str
    pages: tuple[int, ...] = ()
    warnings: tuple[str, ...] = field(default_factory=tuple)

    def metadata(self) -> dict[str, Any]:
        found: dict[str, Any] = {"chars": len(self.text)}
        if self.pages:
            found["pages"] = len(self.pages)
        if self.warnings:
            found["warnings"] = list(self.warnings)
        return found


def _normalize(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ParseError("text documents must be UTF-8") from None


def _pdf(data: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError:
        raise ParseError(f"PDF documents need pypdf: {EXTRA_HINT}") from None
    try:
        reader = PdfReader(io.BytesIO(data))
        pages: list[int] = []
        parts: list[str] = []
        offset = 0
        for page in reader.pages:
            pages.append(offset)
            content = _normalize(page.extract_text() or "").strip()
            parts.append(content)
            offset += len(content) + 2
    except (PdfReadError, ValueError, KeyError) as error:
        raise ParseError(f"unreadable PDF: {error}") from None
    text = "\n\n".join(parts)
    warnings = () if text.strip() else ("no text layer: the PDF looks scanned (OCR is not run)",)
    return ParsedDocument(text, PDF, tuple(pages), warnings)


def _docx(data: bytes) -> ParsedDocument:
    try:
        import docx
    except ImportError:
        raise ParseError(f"DOCX documents need python-docx: {EXTRA_HINT}") from None
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as error:  # python-docx raises zipfile, lxml and its own errors
        raise ParseError(f"unreadable DOCX: {error}") from None
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        if not content:
            continue
        style = (paragraph.style.name if paragraph.style is not None else "") or ""
        if style.startswith("Heading"):
            level = style.removeprefix("Heading").strip()
            depth = int(level) if level.isdigit() else 1
            content = f"{'#' * min(depth, 6)} {content}"
        blocks.append(content)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            blocks.append("| " + " | ".join(cells) + " |")
    return ParsedDocument(_normalize("\n\n".join(blocks)), DOCX)


def parse_document(data: bytes, media_type: str) -> ParsedDocument:
    """Text of a document of a supported media type; raises `ParseError`."""
    if media_type in (MARKDOWN, PLAIN_TEXT):
        return ParsedDocument(_normalize(_decode(data)), media_type)
    if media_type == PDF:
        return _pdf(data)
    if media_type == DOCX:
        return _docx(data)
    raise ParseError(f"unsupported media type {media_type!r}")
